import os
from datetime import datetime
from werkzeug.utils import secure_filename
from flask import Flask, request, jsonify, send_from_directory
from flask_cors import CORS
import json
import uuid
import re

app = Flask(__name__)
CORS(app)

app.config['UPLOAD_FOLDER'] = 'uploads'
app.config['MAX_CONTENT_LENGTH'] = 100 * 1024 * 1024
app.config['HISTORY_FILE'] = 'history.json'

ALLOWED_EXTENSIONS = {'md', 'docx', 'pdf', 'txt'}

if not os.path.exists(app.config['UPLOAD_FOLDER']):
    os.makedirs(app.config['UPLOAD_FOLDER'])

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

def load_history():
    if os.path.exists(app.config['HISTORY_FILE']):
        with open(app.config['HISTORY_FILE'], 'r', encoding='utf-8') as f:
            return json.load(f)
    return []

def save_history(history):
    with open(app.config['HISTORY_FILE'], 'w', encoding='utf-8') as f:
        json.dump(history, f, ensure_ascii=False, indent=2)

@app.route('/api/upload', methods=['POST'])
def upload_file():
    if 'file' not in request.files:
        return jsonify({'error': '没有文件'}), 400

    file = request.files['file']
    if file.filename == '':
        return jsonify({'error': '没有选择文件'}), 400

    if file and allowed_file(file.filename):
        filename = secure_filename(file.filename)
        
        unique_filename = f"{uuid.uuid4().hex}_{filename}"
        filepath = os.path.join(app.config['UPLOAD_FOLDER'], unique_filename)
        file.save(filepath)

        file_info = {
            'id': unique_filename,
            'original_name': filename,
            'stored_name': unique_filename,
            'filepath': filepath,
            'upload_time': datetime.now().isoformat(),
            'status': 'uploaded'
        }

        history = load_history()
        history.append(file_info)
        save_history(history)

        return jsonify({
            'success': True,
            'file': file_info
        })

    return jsonify({'error': '不支持的文件格式'}), 400

@app.route('/api/history', methods=['GET'])
def get_history():
    history = load_history()
    return jsonify({'success': True, 'history': history})

@app.route('/api/parse', methods=['POST'])
def parse_document():
    data = request.json
    file_id = data.get('file_id')

    history = load_history()
    file_info = next((f for f in history if f['id'] == file_id), None)

    if not file_info:
        return jsonify({'error': '文件不存在'}), 404

    filepath = file_info['filepath']
    file_extension = file_info['original_name'].rsplit('.', 1)[1].lower() if '.' in file_info['original_name'] else ''

    try:
        document_data = extract_document_content(filepath, file_extension)
        
        chunks = split_document(document_data['content'])

        return jsonify({
            'success': True,
            'content': document_data['content'],
            'chunks': chunks,
            'tables': document_data.get('tables', []),
            'images': document_data.get('images', []),
            'metadata': document_data.get('metadata', {})
        })
    except Exception as e:
        return jsonify({'error': f'解析失败: {str(e)}'}), 500

def extract_document_content(filepath, file_extension):
    """提取文档内容，包括文本、表格和图片信息"""
    result = {
        'content': '',
        'tables': [],
        'images': [],
        'metadata': {}
    }

    if file_extension == 'md' or file_extension == 'txt':
        with open(filepath, 'r', encoding='utf-8') as f:
            content = f.read()
        result['content'] = content
        result['metadata']['type'] = 'text'
        result['metadata']['line_count'] = len(content.split('\n'))
    
    elif file_extension == 'docx':
        from docx import Document
        
        doc = Document(filepath)
        content_parts = []
        tables = []
        images = []
        
        for element in doc.element.body:
            if element.tag.endswith('p'):
                # 段落
                para = doc.paragraphs[len(content_parts)]
                content_parts.append(para.text)
            elif element.tag.endswith('tbl'):
                # 表格
                table_content = extract_table_content(doc, element)
                if table_content:
                    tables.append(table_content)
                    # 在内容中标记表格位置
                    content_parts.append(f"【表格{len(tables)}】: {table_content['summary']}")
            elif element.tag.endswith('drawing') or element.tag.endswith('pict'):
                # 图片
                image_info = extract_image_info(doc, element)
                images.append(image_info)
                content_parts.append(f"【图片{len(images)}】: {image_info['description']}")
        
        result['content'] = '\n'.join(content_parts)
        result['tables'] = tables
        result['images'] = images
        result['metadata']['type'] = 'docx'
        result['metadata']['paragraph_count'] = len(doc.paragraphs)
        result['metadata']['table_count'] = len(tables)
        result['metadata']['image_count'] = len(images)
    
    elif file_extension == 'pdf':
        from PyPDF2 import PdfReader
        
        reader = PdfReader(filepath)
        content_parts = []
        images = []
        
        for page_num, page in enumerate(reader.pages, 1):
            text = page.extract_text()
            if text:
                content_parts.append(text)
            
            # 检测页面中的图片
            if '/XObject' in page:
                xobjects = page['/XObject'].get_object()
                if xobjects:
                    img_count = sum(1 for key in xobjects if xobjects[key]['/Subtype'] == '/Image')
                    if img_count > 0:
                        images.append({
                            'page': page_num,
                            'description': f'第{page_num}页图片',
                            'count': img_count
                        })
                        content_parts.append(f"【图片{len(images)}】: 第{page_num}页，共{img_count}张图片")
        
        result['content'] = '\n'.join(content_parts)
        result['images'] = images
        result['metadata']['type'] = 'pdf'
        result['metadata']['page_count'] = len(reader.pages)
        result['metadata']['image_count'] = len(images)
    
    else:
        raise ValueError('不支持的文件格式')
    
    return result

def extract_table_content(doc, table_element):
    """从docx表格元素中提取内容"""
    try:
        # 查找表格索引
        table_index = -1
        for i, tbl in enumerate(doc.tables):
            if tbl._element == table_element:
                table_index = i
                break
        
        if table_index == -1:
            return None
        
        table = doc.tables[table_index]
        rows = []
        headers = []
        
        for row_idx, row in enumerate(table.rows):
            cells = []
            for cell in row.cells:
                cells.append(cell.text.strip())
            
            if row_idx == 0:
                headers = cells
                rows.append({'type': 'header', 'cells': cells})
            else:
                rows.append({'type': 'data', 'cells': cells})
        
        # 生成表格摘要
        summary = f"表格（{len(table.rows)}行×{len(table.columns)}列）"
        if headers:
            summary += f"，表头：{', '.join(headers[:3])}{'...' if len(headers) > 3 else ''}"
        
        return {
            'index': table_index + 1,
            'rows': rows,
            'row_count': len(table.rows),
            'col_count': len(table.columns),
            'headers': headers,
            'summary': summary
        }
    except Exception as e:
        return None

def extract_image_info(doc, element):
    """从docx元素中提取图片信息"""
    try:
        # 统计图片数量
        image_count = 0
        for rel in doc.part.rels.values():
            if "image" in rel.target_ref:
                image_count += 1
        
        return {
            'index': image_count,
            'description': '文档图片',
            'type': 'embedded'
        }
    except Exception as e:
        return {
            'index': len(doc.part.rels),
            'description': '文档中的图片',
            'type': 'embedded'
        }

def split_document(content):
    chunks = []
    lines = content.split('\n')
    
    current_section = None
    section_content = []
    
    for line in lines:
        stripped_line = line.strip()
        
        if is_module_title(stripped_line):
            if section_content and current_section:
                chunks.append({
                    'title': current_section,
                    'content': '\n'.join(section_content),
                    'type': 'section'
                })
            
            current_section = stripped_line
            section_content = []
        elif stripped_line and current_section:
            section_content.append(line)
        elif stripped_line:
            if not current_section:
                current_section = '文档概述'
            section_content.append(line)
    
    if section_content:
        chunks.append({
            'title': current_section if current_section else '文档内容',
            'content': '\n'.join(section_content),
            'type': 'section'
        })
    
    if not chunks:
        chunks.append({'title': '文档内容', 'content': content, 'type': 'content'})
    
    return chunks

def is_module_title(line):
    """判断是否为模块标题（支持数字编号格式如 1.、1.1、2. 等）"""
    if not line:
        return False

    # 数字编号标题：1. 标题 / 1.1 标题 / 1.1.1 标题
    if re.match(r'^\d+(?:\.\d+)*[\.、]\s*\S+', line):
        return True

    # 检查 markdown 标题格式
    if line.startswith('#'):
        return True

    return False

@app.route('/api/generate-test-points', methods=['POST'])
def generate_test_points():
    from langchain_openai import ChatOpenAI
    from langchain.prompts import PromptTemplate
    from langchain.chains import LLMChain

    data = request.json
    chunks = data.get('chunks', [])
    tables = data.get('tables', [])
    images = data.get('images', [])

    api_key = os.getenv('OPENAI_API_KEY')
    model_name = data.get('model', 'gpt-3.5-turbo')

    if not api_key:
        test_points = generate_test_points_from_content(chunks, tables, images)
        return jsonify({
            'error': '未配置API密钥，使用基于文档内容的测试点生成',
            'test_points': test_points
        }), 200

    try:
        llm = ChatOpenAI(
            api_key=api_key,
            model=model_name,
            temperature=0.7
        )

        prompt_template = PromptTemplate(
            input_variables=["content"],
            template="""你是一位资深软件测试工程师。请分析以下需求文档内容，提取测试点。

要求：
1. 输出JSON格式的测试点列表
2. 每个测试点包含：test_point_id, category, description, priority, test_type
3. test_type可选值：功能测试, 交互测试, 异常测试, 边界测试, 权限测试, 性能测试
4. priority可选值：高, 中, 低
5. category分类：功能点, 交互点, 异常点, 边界点, 校验点

需求文档内容：
{content}

请生成测试点列表（JSON格式）："""
        )

        chain = LLMChain(llm=llm, prompt=prompt_template)
        all_test_points = []

        for idx, chunk in enumerate(chunks):
            response = chain.invoke({"content": chunk['content'][:2000]})
            response_text = normalize_llm_response(response)

            try:
                points = parse_llm_points(response_text)
                if isinstance(points, list):
                    for point in points:
                        normalize_test_point(point, idx, len(all_test_points) + 1)
                    all_test_points.extend(points)
            except json.JSONDecodeError:
                points = parse_text_test_points(response_text)
                for point in points:
                    normalize_test_point(point, idx, len(all_test_points) + 1)
                all_test_points.extend(points)

        # 对AI结果做质量兜底：低质量结果自动使用规则引擎补全
        all_test_points = improve_generated_points(all_test_points, chunks, tables, images)

        return jsonify({
            'success': True,
            'test_points': all_test_points
        })

    except Exception as e:
        return jsonify({
            'error': f'AI生成失败: {str(e)}',
            'test_points': generate_test_points_from_content(chunks, tables, images)
        }), 200

def parse_text_test_points(text):
    points = []
    lines = text.split('\n')
    current_point = None

    for line in lines:
        line = line.strip()
        if not line or line.startswith('#') or line.startswith('```'):
            continue

        if line and (line[0].isdigit() or line.startswith('-') or line.startswith('*')):
            if current_point:
                points.append(current_point)

            current_point = {
                'test_point_id': f"TP-{len(points) + 1:03d}",
                'category': '功能点',
                'description': line.lstrip('0123456789.-* ').strip(),
                'priority': '中',
                'test_type': '功能测试'
            }

    if current_point:
        points.append(current_point)

    return points

def normalize_llm_response(response):
    """兼容 LangChain 不同返回结构，统一提取字符串内容"""
    if isinstance(response, str):
        return response
    if isinstance(response, dict):
        for key in ('text', 'output_text', 'content'):
            value = response.get(key)
            if isinstance(value, str):
                return value
        return json.dumps(response, ensure_ascii=False)
    return str(response)

def parse_llm_points(response_text):
    """解析大模型返回的测试点 JSON（兼容 Markdown 代码块）"""
    if not response_text:
        return []

    cleaned = response_text.strip()
    if cleaned.startswith('```'):
        cleaned = cleaned.strip('`')
        if cleaned.lower().startswith('json'):
            cleaned = cleaned[4:]
        cleaned = cleaned.strip()

    parsed = json.loads(cleaned)
    if isinstance(parsed, dict):
        if isinstance(parsed.get('test_points'), list):
            return parsed['test_points']
        return [parsed]
    return parsed

def normalize_test_point(point, chunk_id, index):
    """补齐测试点必填字段，避免模型漏字段导致前端异常"""
    point['chunk_id'] = chunk_id
    point['test_point_id'] = f"TP-{index:03d}"
    point.setdefault('category', '功能点')
    point.setdefault('description', '待补充测试点描述')
    point.setdefault('priority', '中')
    point.setdefault('test_type', '功能测试')

def is_low_quality_point(point):
    """识别低质量测试点：描述过泛、模板痕迹重、缺少业务可执行性"""
    description = str(point.get('description', '')).strip()
    title = str(point.get('title', '')).strip()
    noisy_keywords = [
        '规则动作[', '字段[字段名称', '存在性检查', '正确性',
        '准备测试数据', '执行触发规则', '验证规则执行结果'
    ]
    if len(description) < 10:
        return True
    if any(word in description for word in noisy_keywords):
        return True
    if title.startswith('规则规则') or title.endswith('正确性'):
        return True
    return False

def improve_generated_points(points, chunks, tables=None, images=None):
    """保留高质量AI结果，并用规则提取结果补齐关键业务测试点"""
    high_quality_points = [p for p in points if not is_low_quality_point(p)]
    if len(high_quality_points) >= 20:
        return reindex_test_points(high_quality_points)

    # AI输出质量不足时，切换到规则引擎（更稳定、可解释）
    rule_points = generate_test_points_from_content(chunks, tables, images)
    return reindex_test_points(high_quality_points + rule_points)

def reindex_test_points(points):
    seen = set()
    unique_points = []
    for point in points:
        dedupe_key = (
            point.get('module', ''),
            point.get('sub_module', ''),
            point.get('description', '')
        )
        if dedupe_key in seen:
            continue
        seen.add(dedupe_key)
        unique_points.append(point)

    for idx, point in enumerate(unique_points, 1):
        point['test_point_id'] = f"TP-{idx:03d}"
    return unique_points

def generate_demo_test_points():
    return [
        {
            'test_point_id': 'TP-001',
            'category': '功能点',
            'description': '验证用户能够成功上传需求文档',
            'priority': '高',
            'test_type': '功能测试',
            'chunk_id': 0
        },
        {
            'test_point_id': 'TP-002',
            'category': '功能点',
            'description': '验证系统能够正确解析文档内容',
            'priority': '高',
            'test_type': '功能测试',
            'chunk_id': 0
        },
        {
            'test_point_id': 'TP-003',
            'category': '异常点',
            'description': '验证上传不支持的文件格式时的错误处理',
            'priority': '中',
            'test_type': '异常测试',
            'chunk_id': 0
        },
        {
            'test_point_id': 'TP-004',
            'category': '边界点',
            'description': '验证超大文档的分块处理能力',
            'priority': '中',
            'test_type': '边界测试',
            'chunk_id': 0
        },
        {
            'test_point_id': 'TP-005',
            'category': '校验点',
            'description': '验证文档内容提取的完整性',
            'priority': '高',
            'test_type': '功能测试',
            'chunk_id': 0
        },
        {
            'test_point_id': 'TP-006',
            'category': '交互点',
            'description': '验证用户界面操作的流畅性',
            'priority': '中',
            'test_type': '交互测试',
            'chunk_id': 1
        },
        {
            'test_point_id': 'TP-007',
            'category': '功能点',
            'description': '验证测试点的手动编辑功能',
            'priority': '高',
            'test_type': '功能测试',
            'chunk_id': 1
        },
        {
            'test_point_id': 'TP-008',
            'category': '异常点',
            'description': '验证空文档上传的处理',
            'priority': '低',
            'test_type': '异常测试',
            'chunk_id': 1
        },
        {
            'test_point_id': 'TP-009',
            'category': '权限测试',
            'description': '验证未授权用户访问控制',
            'priority': '高',
            'test_type': '权限测试',
            'chunk_id': 2
        },
        {
            'test_point_id': 'TP-010',
            'category': '性能测试',
            'description': '验证大量测试点生成时的响应时间',
            'priority': '中',
            'test_type': '性能测试',
            'chunk_id': 2
        }
    ]

def generate_test_points_from_content(chunks, tables=None, images=None):
    """基于文档规则提取测试点（结构化、可执行、去模板化）"""
    tables = tables or []
    images = images or []
    all_lines = []
    for chunk in chunks:
        all_lines.extend(chunk.get('content', '').split('\n'))

    extracted = extract_structured_points_from_lines(all_lines)

    # 文档有表格/图片时补充解析质量校验点
    if tables:
        extracted.append(build_point(
            module='文档解析',
            sub_module='表格解析',
            category='校验点',
            description='验证需求文档中表格字段均被完整提取且语义不丢失',
            priority='高',
            test_type='功能测试'
        ))
    if images:
        extracted.append(build_point(
            module='文档解析',
            sub_module='图片解析',
            category='校验点',
            description='验证文档中的图片说明可被识别并参与测试点推导',
            priority='中',
            test_type='功能测试'
        ))

    return reindex_test_points(extracted)

def extract_structured_points_from_lines(lines):
    module = '文档功能'
    sub_module = '通用'
    points = []
    for raw in lines:
        line = raw.strip()
        if not line:
            continue
        if line.startswith('### '):
            sub_module = line.replace('###', '').strip()
            continue
        if line.startswith('## '):
            module = line.replace('##', '').strip()
            continue
        if line.startswith('#### '):
            sub_module = line.replace('####', '').strip()
            continue

        # 通用表格行解析：| 字段 | 类型 | 是否必填 | 说明 |
        if is_table_data_line(line):
            fields = parse_markdown_table_row(line)
            if fields:
                field_name = fields[0]
                if field_name and field_name not in ('字段名称', '字段', '参数'):
                    points.append(build_point(
                        module, sub_module, '校验点',
                        f'验证字段[{field_name}]在页面可见且满足文档定义',
                        '中', '功能测试'
                    ))
            continue

        if should_skip_line(line):
            continue

        # 公式类测试点
        if '=' in line and any(k in line for k in ['成本', '费用', '总成本', '公式']):
            points.append(build_point(
                module, sub_module, '校验点',
                f'验证{sub_module}公式计算与文档定义一致：{safe_clip(line, 50)}',
                '高', '功能测试'
            ))
            continue

        # 强规则类：必须/不允许/仅支持/至少/最多/失败时
        if any(k in line for k in ['必须', '不允许', '仅支持', '至少', '最多', '失败时', '阻断', '不可']):
            points.append(build_point(
                module, sub_module, '异常点',
                f'验证规则约束生效：{safe_clip(line, 50)}',
                '高', '异常测试'
            ))
            continue

        # 业务动作类：支持/导入/导出/新增/编辑/删除/查询/同步/匹配
        if any(k in line for k in ['支持', '导入', '导出', '新增', '编辑', '删除', '查询', '同步', '匹配', '核算']):
            points.append(build_point(
                module, sub_module, '功能点',
                f'验证业务流程：{safe_clip(line, 50)}',
                '中', '功能测试'
            ))
            continue

    # 兜底：确保至少生成通用关键场景
    if len(points) < 10:
        points.extend(build_default_generic_points(module, sub_module))
    return points

def should_skip_line(line):
    if re.match(r'^[-*]{3,}$', line):
        return True
    if line.startswith('```'):
        return True
    if len(line) <= 4:
        return True
    return False

def safe_clip(text, max_len=50):
    clipped = text.replace('**', '').replace('`', '').strip()
    return clipped[:max_len]

def build_point(module, sub_module, category, description, priority, test_type):
    return {
        'module': module,
        'sub_module': sub_module,
        'category': category,
        'description': description,
        'priority': priority,
        'test_type': test_type,
        'chunk_id': 0
    }

def is_table_data_line(line):
    return line.startswith('|') and line.endswith('|') and line.count('|') >= 2

def parse_markdown_table_row(line):
    cells = [cell.strip() for cell in line.strip('|').split('|')]
    if not cells:
        return []
    # 排除分隔行，如 |----|----|
    if all(re.match(r'^:?-{2,}:?$', c) for c in cells if c):
        return []
    return cells

def build_default_generic_points(module, sub_module):
    """通用兜底测试点（与具体行业无关）"""
    return [
        build_point(module, sub_module, '功能点', '验证主流程可以从输入到结果完整闭环执行', '高', '功能测试'),
        build_point(module, sub_module, '异常点', '验证必填项缺失时系统可拦截并提示明确错误', '高', '异常测试'),
        build_point(module, sub_module, '边界点', '验证边界值输入时系统行为符合预期', '中', '边界测试'),
        build_point(module, sub_module, '校验点', '验证关键计算或规则判定结果与文档定义一致', '高', '功能测试'),
        build_point(module, sub_module, '交互点', '验证页面按钮与状态切换交互一致且可追踪', '中', '交互测试'),
        build_point(module, sub_module, '异常点', '验证外部依赖异常时系统有降级和错误提示', '中', '异常测试'),
        build_point(module, sub_module, '功能点', '验证批量操作支持部分成功并输出失败明细', '中', '功能测试'),
        build_point(module, sub_module, '校验点', '验证导入导出字段映射完整且顺序正确', '中', '功能测试'),
    ]

def extract_modules_from_chunks(chunks):
    """从分块中提取模块"""
    modules = []
    for chunk in chunks:
        title = chunk.get('title', '')
        content = chunk.get('content', '')
        
        if title and title != '文档整体' and title != '文档内容' and title != '文档概述':
            modules.append({
                'name': title,
                'content': content
            })
    return modules

def parse_modules_from_content(content):
    """从文档内容中解析模块（支持数字编号格式）"""
    modules = []
    lines = content.split('\n')
    
    current_module = None
    current_content = []
    
    for line in lines:
        line = line.strip()
        
        # 匹配模块标题格式：数字. 标题（如 "1. 游客可见商品清单"）
        if is_module_title(line):
            # 保存上一个模块
            if current_module:
                modules.append({
                    'name': current_module,
                    'content': '\n'.join(current_content)
                })
            
            # 提取新模块名称
            current_module = line
            current_content = []
        elif current_module:
            current_content.append(line)
    
    # 保存最后一个模块
    if current_module:
        modules.append({
            'name': current_module,
            'content': '\n'.join(current_content)
        })
    
    # 如果没有解析到模块，返回一个默认模块
    if not modules:
        modules.append({
            'name': '文档功能',
            'content': content
        })
    
    return modules

def clean_module_name_func(module_name):
    """清理模块名称（移除编号）"""
    if not module_name:
        return '功能模块'
    
    # 移除开头的数字编号
    if module_name[0].isdigit():
        parts = module_name.split('.', 1)
        if len(parts) > 1:
            return parts[1].strip()
    
    # 移除 markdown 标题符号
    return module_name.strip('# ')

@app.route('/api/generate-test-cases', methods=['POST'])
def generate_test_cases():
    from langchain_openai import ChatOpenAI
    from langchain.prompts import PromptTemplate
    from langchain.chains import LLMChain

    data = request.json
    test_points = data.get('test_points', [])

    api_key = os.getenv('OPENAI_API_KEY')
    model_name = data.get('model', 'gpt-3.5-turbo')

    if not api_key:
        return jsonify({
            'error': '未配置API密钥',
            'test_cases': generate_demo_test_cases(test_points)
        }), 200

    try:
        llm = ChatOpenAI(
            api_key=api_key,
            model=model_name,
            temperature=0.7
        )

        prompt_template = PromptTemplate(
            input_variables=["test_point"],
            template="""你是一位资深软件测试工程师。请基于以下测试点生成标准测试用例，严格遵循BlueTool测试用例生成规范。

测试点信息：
{test_point}

要求：
1. 输出JSON格式的测试用例
2. 每个测试用例包含：case_id, title, module, preconditions, steps, expected_results, priority, involve_side, scene
3. case_id格式：TC-XXX（序号从1开始）
4. priority可选值：高, 中, 低
5. steps必须是具体的操作步骤，每步一行，格式为"1. 操作描述"
6. expected_results必须与steps一一对应，每步一行，格式为"1. 预期结果"
7. involve_side固定为"Web端"
8. scene根据测试类型填写：正向/异常校验/边界值/权限控制/网络异常
9. 测试用例必须可执行、无模糊表述
10. 覆盖正向、反向全场景

请生成测试用例（JSON格式）："""
        )

        chain = LLMChain(llm=llm, prompt=prompt_template)
        all_cases = []
        test_point_chunks = {}

        for point in test_points:
            chunk_id = point.get('chunk_id', 0)
            if chunk_id not in test_point_chunks:
                test_point_chunks[chunk_id] = []
            test_point_chunks[chunk_id].append(point)

        for chunk_id, points in test_point_chunks.items():
            combined_content = '\n'.join([
                f"{p.get('test_point_id', '')}: {p.get('description', '')} (类型: {p.get('test_type', '功能测试')}, 优先级: {p.get('priority', '中')})"
                for p in points
            ])

            try:
                response = chain.invoke({"test_point": combined_content})
                cases = json.loads(response)

                if isinstance(cases, list):
                    for case in cases:
                        case['module'] = f"模块{chunk_id + 1}"
                        case['case_id'] = f"TC-{len(all_cases) + 1:03d}"
                        case['related_test_points'] = [p.get('test_point_id', '') for p in points]
                    all_cases.extend(cases)
                elif isinstance(cases, dict):
                    cases['module'] = f"模块{chunk_id + 1}"
                    cases['case_id'] = f"TC-{len(all_cases) + 1:03d}"
                    cases['related_test_points'] = [p.get('test_point_id', '') for p in points]
                    all_cases.append(cases)
            except (json.JSONDecodeError, Exception) as e:
                for point in points:
                    case = generate_single_case_from_point(point, len(all_cases) + 1)
                    all_cases.append(case)

        return jsonify({
            'success': True,
            'test_cases': all_cases
        })

    except Exception as e:
        return jsonify({
            'error': f'AI生成失败: {str(e)}',
            'test_cases': generate_demo_test_cases(test_points)
        }), 200

def generate_single_case_from_point(point, index):
    description = point.get('description', '')
    test_type = point.get('test_type', '功能测试')
    priority = point.get('priority', '中')
    
    # 智能提取模块名
    module_name = extract_module_name(description)
    
    # 生成简洁的测试用例标题
    title = generate_case_title(description, test_type)
    
    # 生成贴合实际的前置条件
    preconditions = generate_preconditions(module_name, test_type)
    
    # 生成具体的测试步骤和预期结果
    steps, expected_results = generate_steps_and_results(description, module_name, test_type)
    
    # 确定场景类型
    scene = get_scene_type(test_type)
    
    return {
        'case_id': f"TC-{index:03d}",
        'title': title,
        'module': module_name,
        'preconditions': preconditions,
        'steps': steps,
        'expected_results': expected_results,
        'priority': priority,
        'related_test_points': [point.get('test_point_id', '')],
        'involve_side': 'Web端',
        'scene': scene
    }

def extract_module_name(description):
    """从测试点描述中智能提取模块名称"""
    if not description:
        return '功能模块'
    
    # 移除开头的"验证"
    if description.startswith('验证'):
        remaining = description[2:]
    else:
        remaining = description
    
    # 提取第一个冒号或"的"之前的内容作为模块名
    separators = ['：', ':', '的', '——', '-']
    for sep in separators:
        if sep in remaining:
            module_part = remaining.split(sep)[0].strip()
            if module_part and len(module_part) <= 20:
                return module_part
    
    # 如果没有找到分隔符，取前20个字符
    return remaining[:20].strip()

def generate_case_title(description, test_type):
    """生成简洁的测试用例标题（不超过30字）"""
    # 移除开头的"验证"
    title_base = description.replace('验证', '').strip()
    
    # 根据测试类型添加前缀
    type_prefix = {
        '功能测试': '功能',
        '异常测试': '异常',
        '边界测试': '边界',
        '交互测试': '交互',
        '权限测试': '权限',
        '性能测试': '性能'
    }
    
    prefix = type_prefix.get(test_type, '功能')
    
    # 生成简洁标题
    full_title = f"{prefix}-{title_base}"
    
    # 截断到30字
    if len(full_title) > 30:
        full_title = full_title[:27] + '...'
    
    return full_title

def generate_preconditions(module_name, test_type):
    """根据模块和测试类型生成贴合实际的前置条件"""
    base_conditions = [
        '系统已启动并正常运行',
        '测试环境已准备就绪',
        f'{module_name}功能页面可正常访问'
    ]
    
    # 根据测试类型添加特定前置条件
    type_specific = {
        '功能测试': [],
        '异常测试': ['准备异常测试数据（如无效格式、空值等）'],
        '边界测试': ['准备边界值测试数据（如最大/最小值等）'],
        '交互测试': ['确保网络连接稳定'],
        '权限测试': ['已准备不同权限的测试账号'],
        '性能测试': ['确保测试环境资源充足']
    }
    
    additional = type_specific.get(test_type, [])
    
    return base_conditions + additional

def generate_steps_and_results(description, module_name, test_type):
    """生成具体的测试步骤和预期结果"""
    # 移除"验证"并获取核心操作描述
    core_action = description.replace('验证', '').replace('的核心功能', '').replace('的异常输入处理', '').replace('的边界条件处理', '').strip()
    
    if test_type == '功能测试':
        steps = [
            f"1. 进入【{module_name}】功能页面",
            f"2. 执行：{core_action}",
            f"3. 确认操作结果",
            f"4. 验证功能完整性"
        ]
        expected_results = [
            f"1. 成功进入{module_name}页面，界面加载正常",
            f"2. 操作执行成功，无报错信息",
            f"3. 操作结果与预期一致",
            f"4. {module_name}功能完整可用"
        ]
    
    elif test_type == '异常测试':
        steps = [
            f"1. 进入【{module_name}】功能页面",
            f"2. 输入异常数据或执行异常操作",
            f"3. 观察系统响应",
            f"4. 验证错误处理结果"
        ]
        expected_results = [
            f"1. 成功进入{module_name}页面",
            f"2. 系统接收异常输入",
            f"3. 系统显示友好错误提示，提示内容清晰明确",
            f"4. 系统不崩溃，可继续操作"
        ]
    
    elif test_type == '边界测试':
        steps = [
            f"1. 进入【{module_name}】功能页面",
            f"2. 输入边界值数据进行测试",
            f"3. 观察系统响应",
            f"4. 验证边界处理能力"
        ]
        expected_results = [
            f"1. 成功进入{module_name}页面",
            f"2. 边界值数据被正确接收",
            f"3. 系统响应符合预期",
            f"4. 边界条件处理正确，无数据溢出或截断问题"
        ]
    
    elif test_type == '交互测试':
        steps = [
            f"1. 进入【{module_name}】功能页面",
            f"2. 执行用户交互操作（点击、输入等）",
            f"3. 观察界面响应和状态变化",
            f"4. 验证交互流畅性"
        ]
        expected_results = [
            f"1. 成功进入{module_name}页面",
            f"2. 交互操作响应及时",
            f"3. 界面状态正确更新，无卡顿",
            f"4. 用户体验流畅，无操作延迟"
        ]
    
    elif test_type == '权限测试':
        steps = [
            f"1. 登录测试账号",
            f"2. 尝试访问【{module_name}】功能",
            f"3. 验证操作权限",
            f"4. 切换不同权限账号重复测试"
        ]
        expected_results = [
            f"1. 成功登录系统",
            f"2. 根据权限显示/隐藏功能入口",
            f"3. 权限范围内操作正常，越权操作被拦截",
            f"4. 权限控制逻辑正确"
        ]
    
    else:  # 性能测试
        steps = [
            f"1. 准备性能测试环境",
            f"2. 对【{module_name}】功能执行压力测试",
            f"3. 记录响应时间和资源消耗",
            f"4. 分析性能指标"
        ]
        expected_results = [
            f"1. 测试环境准备完成",
            f"2. 压力测试执行成功",
            f"3. 响应时间在可接受范围内，资源消耗正常",
            f"4. 性能指标符合要求"
        ]
    
    return steps, expected_results

def get_scene_type(test_type):
    """根据测试类型确定场景"""
    scene_map = {
        '功能测试': '正向',
        '异常测试': '异常校验',
        '边界测试': '边界值',
        '交互测试': '正向',
        '权限测试': '权限控制',
        '性能测试': '正向'
    }
    return scene_map.get(test_type, '正向')

def generate_demo_test_cases(test_points):
    cases = []
    for idx, point in enumerate(test_points, 1):
        case = generate_single_case_from_point(point, idx)
        cases.append(case)
    return cases

@app.route('/api/export', methods=['POST'])
def export_cases():
    data = request.json
    test_cases = data.get('test_cases', [])
    format_type = data.get('format', 'json')

    if format_type == 'json':
        return jsonify({
            'success': True,
            'data': json.dumps(test_cases, ensure_ascii=False, indent=2)
        })

    elif format_type == 'markdown':
        md_content = "# 测试用例文档\n\n"
        md_content += f"生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n\n"

        md_content += "## 标准化测试用例库\n\n"
        md_content += "| 所属模块 | 测试标题 | 涉及端侧 | 场景 | 前置条件 | 测试步骤 | 预期结果 | 优先级 |\n"
        md_content += "| :--- | :--- | :--- | :--- | :--- | :--- | :--- | :--- |\n"

        for case in test_cases:
            steps = '\\n'.join(case.get('steps', []))
            expected_results = '\\n'.join(case.get('expected_results', []))
            preconditions = '\\n'.join(case.get('preconditions', []))
            
            md_content += f"| {case.get('module', '未分类')} | {case.get('case_id', 'TC-000')}: {case.get('title', '未命名')} | {case.get('involve_side', 'Web端')} | {case.get('scene', '正向')} | {preconditions} | {steps} | {expected_results} | {case.get('priority', '中')} |\n"
        
        md_content += "\n---\n\n"
        
        md_content += "## 详细测试用例\n\n"
        for case in test_cases:
            md_content += f"### {case.get('case_id', 'TC-000')}: {case.get('title', '未命名')}\n\n"
            md_content += f"**所属模块**: {case.get('module', '未分类')}\n"
            md_content += f"**涉及端侧**: {case.get('involve_side', 'Web端')}\n"
            md_content += f"**场景**: {case.get('scene', '正向')}\n"
            md_content += f"**优先级**: {case.get('priority', '中')}\n\n"
            md_content += f"**前置条件**:\n"
            for pre in case.get('preconditions', []):
                md_content += f"- {pre}\n"
            md_content += "\n"

            md_content += f"**测试步骤**:\n"
            for i, step in enumerate(case.get('steps', []), 1):
                md_content += f"{i}. {step}\n"
            md_content += "\n"

            md_content += f"**预期结果**:\n"
            for i, result in enumerate(case.get('expected_results', []), 1):
                md_content += f"{i}. {result}\n"
            md_content += "\n---\n\n"

        return jsonify({
            'success': True,
            'data': md_content
        })
    
    elif format_type == 'excel':
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
            from openpyxl.utils import get_column_letter
            
            wb = Workbook()
            ws = wb.active
            ws.title = "测试用例"
            
            # 设置表头样式
            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="409EFF", end_color="409EFF", fill_type="solid")
            header_alignment = Alignment(horizontal="center", vertical="center")
            thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), 
                               top=Side(style='thin'), bottom=Side(style='thin'))
            
            # 定义表头
            headers = ['用例编号', '测试标题', '所属模块', '涉及端侧', '场景', '优先级', 
                       '前置条件', '测试步骤', '预期结果']
            
            # 写入表头
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
                cell.border = thin_border
            
            # 设置列宽
            column_widths = [12, 30, 15, 10, 12, 8, 30, 40, 40]
            for i, width in enumerate(column_widths, 1):
                ws.column_dimensions[get_column_letter(i)].width = width
            
            # 写入测试用例数据
            data_fill = PatternFill(start_color="F5F7FA", end_color="F5F7FA", fill_type="solid")
            
            for row, case in enumerate(test_cases, 2):
                ws.cell(row=row, column=1, value=clean_string(case.get('case_id', ''))).border = thin_border
                ws.cell(row=row, column=2, value=clean_string(case.get('title', ''))).border = thin_border
                ws.cell(row=row, column=3, value=clean_string(case.get('module', ''))).border = thin_border
                ws.cell(row=row, column=4, value=clean_string(case.get('involve_side', 'Web端'))).border = thin_border
                ws.cell(row=row, column=5, value=clean_string(case.get('scene', ''))).border = thin_border
                ws.cell(row=row, column=6, value=clean_string(case.get('priority', ''))).border = thin_border
                
                preconditions = '\n'.join([clean_string(p) for p in case.get('preconditions', [])])
                ws.cell(row=row, column=7, value=preconditions).border = thin_border
                
                steps = '\n'.join([clean_string(s) for s in case.get('steps', [])])
                ws.cell(row=row, column=8, value=steps).border = thin_border
                
                expected_results = '\n'.join([clean_string(e) for e in case.get('expected_results', [])])
                ws.cell(row=row, column=9, value=expected_results).border = thin_border
            
            # 保存文件
            excel_filename = f"test_cases_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            excel_path = os.path.join(app.config['UPLOAD_FOLDER'], excel_filename)
            wb.save(excel_path)
            
            return jsonify({
                'success': True,
                'data': excel_filename,
                'message': 'Excel文件已生成'
            })
        
        except Exception as e:
            return jsonify({'error': f'Excel导出失败: {str(e)}'}), 500

    return jsonify({'error': '不支持的导出格式'}), 400

def clean_string(s):
    """清理字符串中的特殊字符，使其可以用于Excel"""
    if not s:
        return ''
    
    # 移除或替换Excel不支持的字符
    s = str(s)
    
    # 移除控制字符（除了换行符）
    cleaned = []
    for char in s:
        # 保留常见的可打印字符和换行符
        if ord(char) < 32 and char != '\n' and char != '\t':
            continue
        # 移除特殊Unicode字符
        if ord(char) > 65535:
            continue
        cleaned.append(char)
    
    return ''.join(cleaned)

@app.route('/api/download-excel/<filename>', methods=['GET'])
def download_excel(filename):
    return send_from_directory(app.config['UPLOAD_FOLDER'], filename, as_attachment=True)

if __name__ == '__main__':
    app.run(debug=True, port=5000, host='0.0.0.0')
